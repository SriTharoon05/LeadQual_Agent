"""Regenerate the two Word guides using repository policy and tariff defaults.

From the repository root, with backend dependencies and python-docx installed:
    backend/.venv/bin/python docs/generate_documents.py
No credentials, network calls, or live customer records are required.
"""
import ast
from decimal import Decimal
from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
from app.scoring.policy import SCORING_VERSION, WEIGHTS, qualification
from app.validators.evidence import FIELDS, GROWTH_STRENGTH

PRICING = next(ast.literal_eval(node.value) for node in ast.parse(
    (ROOT / "backend/app/admin/settings.py").read_text()).body
    if isinstance(node, ast.Assign) and any(isinstance(target, ast.Name)
        and target.id == "DEFAULT_PRICING" for target in node.targets))
DATE = "27 October 2026"
TEAL, INK, LIGHT = "126A76", "183349", "EDF5F7"


def shade(cell, fill):
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(element)


def field(paragraph, name):
    element = OxmlElement("w:fldSimple")
    element.set(qn("w:instr"), name)
    paragraph._p.append(element)


def document(title, subtitle):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.7)
    section.header_distance = section.footer_distance = Inches(.3)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in (("Title", 28), ("Heading 1", 19), ("Heading 2", 13)):
        style = doc.styles[name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor.from_string(TEAL)
    header = section.header.paragraphs[0]
    header.text = "QUALITY LEAD  /  PRODUCT & OPERATIONS"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string(TEAL)
    footer = section.footer.paragraphs[0]
    footer.text = f"{DATE}  •  {title}    |    "
    footer.runs[0].font.size = Pt(8)
    field(footer, "PAGE")
    footer.add_run(" / ")
    field(footer, "NUMPAGES")
    doc.core_properties.title = title
    doc.core_properties.subject = subtitle
    doc.core_properties.author = "Quality Lead"
    doc.add_heading(title, 0)
    doc.add_paragraph(subtitle, style="Subtitle")
    doc.add_paragraph(f"Version 1.0  •  {DATE}  •  Current dev implementation")
    return doc


def para(doc, text):
    doc.add_paragraph(text)


def heading(doc, text):
    doc.add_heading(text, 2)


def page(doc, number, title):
    doc.add_page_break()
    doc.add_heading(f"{number:02d}  {title}", 1)


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths is None:
        widths = [6.87 / len(headers)] * len(headers)
    for column, width in zip(t.columns, widths):
        column.width = Inches(width)
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for cell, label in zip(t.rows[0].cells, headers):
        cell.text = str(label)
        shade(cell, TEAL)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for index, values in enumerate(rows):
        cells = t.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = str(value)
            if index % 2 == 0:
                shade(cell, LIGHT)
    for row in t.rows:
        keep = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(keep)
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def money(amount):
    return f"${Decimal(str(amount)):,.4f}"


def cost_guide():
    doc = document("Costing Guide", "AI usage, budgeting assumptions and operating costs")
    para(doc, "This guide explains what Quality Lead meters, how the application estimates charges, "
         "and how to plan a batch budget. Reference tariffs are repository defaults, not verified Azure "
         "contract prices. No production usage, billing account or live exchange-rate data was accessed.")
    heading(doc, "01  Reference price sheet (USD)")
    table(doc, ["Meter", "Reference rate", "Billing basis"], [
        ["GPT-4.1 mini input", f"${PRICING['input_per_million']:.2f}", "Per 1,000,000 uncached input tokens"],
        ["GPT-4.1 mini cached input", f"${PRICING['cached_per_million']:.2f}", "Per 1,000,000 provider-reported cached tokens"],
        ["GPT-4.1 mini output", f"${PRICING['output_per_million']:.2f}", "Per 1,000,000 output tokens"],
        ["text-embedding-ada-002", f"${PRICING['embedding_per_million']:.2f}", "Per 1,000,000 embedding input tokens"],
        ["Bing grounding / search", f"${PRICING['search_per_thousand']:.2f}", "Per 1,000 reported search requests"],
    ], [2.05, 1.22, 3.6])
    heading(doc, "Before approving a budget")
    para(doc, "Confirm region, deployment type, model version and commercial terms in Azure. "
         "The default pricing setting is unconfirmed. Admin Settings supports editable rates and "
         "a USD/INR reference rate. Rates in this document exclude tax, infrastructure and support.")
    table(doc, ["Term", "Meaning"], [
        ["Token", "A provider-reported input/output unit; not a word or a lead."],
        ["Cached input", "A subset of total input, charged at the cached rate rather than twice."],
        ["Search request", "Provider tool_usage.web_search.num_requests; not the number of visible tool items."],
        ["Cost estimate", "Reported usage multiplied by a saved tariff; not an invoice or guaranteed final charge."],
    ], [1.45, 5.42])

    page(doc, 2, "Calculation and worked example")
    table(doc, ["Component", "USD formula"], [
        ["Uncached input", "(Input tokens − cached input tokens) × input rate / 1,000,000"],
        ["Cached input", "Cached input tokens × cached rate / 1,000,000"],
        ["Output", "Output tokens × output rate / 1,000,000"],
        ["Embeddings", "Embedding input tokens × embedding rate / 1,000,000"],
        ["Search", "Provider-reported search requests × search rate / 1,000"],
        ["Total", "Sum of priced components; any missing meter means an incomplete estimate."],
        ["INR display", "USD estimate × configured INR per USD; reference conversion only."],
    ], [1.4, 5.47])
    heading(doc, "Illustrative one-lead budget")
    para(doc, "Assume aggregate usage across a lead's requests: 12,000 input tokens including 3,000 "
         "cached tokens; 2,000 output tokens; 2,000 embedding tokens; and 3 reported search requests. "
         "These are planning assumptions, not measured averages. Import and chat are excluded.")
    components = [
        ("Uncached input", "9,000", Decimal(9000) * Decimal(str(PRICING['input_per_million'])) / 1000000),
        ("Cached input", "3,000", Decimal(3000) * Decimal(str(PRICING['cached_per_million'])) / 1000000),
        ("Output", "2,000", Decimal(2000) * Decimal(str(PRICING['output_per_million'])) / 1000000),
        ("Embedding", "2,000", Decimal(2000) * Decimal(str(PRICING['embedding_per_million'])) / 1000000),
        ("Search", "3 requests", Decimal(3) * Decimal(str(PRICING['search_per_thousand'])) / 1000),
    ]
    unit = sum(row[2] for row in components)
    table(doc, ["Component", "Quantity", "Estimated USD"],
          [[label, quantity, money(value)] for label, quantity, value in components]
          + [["Total per newly researched lead", "Assumed workload", money(unit)]], [2.9, 1.8, 2.17])
    para(doc, "At an illustrative 85 INR per USD (not a current exchange-rate quote), this is "
         f"INR {unit * 85:.4f} per lead. Keep full precision in calculations and round for display only.")

    page(doc, 3, "Batch budgets and research reuse")
    para(doc, "The following scenarios keep token usage constant at the assumptions on page 2. "
         "Only search requests vary. All leads are newly researched; no retries, import or chat are included.")
    token_cost = unit - components[-1][2]
    table(doc, ["New leads", "1 search / lead", "2 searches / lead", "3 searches / lead"], [
        [f"{n:,}", *[money(n * (token_cost + Decimal(searches) * Decimal(str(PRICING['search_per_thousand'])) / 1000))
                     for searches in (1, 2, 3)]] for n in (100, 1000, 10000)
    ], [1.2, 1.89, 1.89, 1.89])
    heading(doc, "What each processing mode changes")
    table(doc, ["Path", "Usage that can occur", "Budget implication"], [
        ["Fresh AI research", "New web research, ICP comparison, optional outreach draft and embeddings.", "Use the full new-lead assumption; retries may add cost."],
        ["Full result reuse", "Compatible saved scores/results copied with no new per-lead qualification AI calls.", "Zero new qualification AI cost for that copy; import/hosting still apply."],
        ["Research-only reuse", "Web research skipped; contact/ICP/scoring still run; job-scoped indexing may occur.", "Lower expected search cost; not equivalent to a free full result reuse."],
        ["Resume after pause", "New AI calls and indexing wait while paused; successful saved stages are restored on resume.", "In-flight calls and late usage can still be billed. Interrupted embedding documents may need batches recomputed on resume."],
        ["Manual evidence refresh", "Fresh research and recalculation for a completed/failed lead.", "A separate operation; budget separately from resuming a job."],
    ], [1.35, 2.95, 2.57])
    para(doc, f"Mixed-batch example: 1,000 leads with 600 full reuses and 400 newly researched leads "
         f"at {money(unit)} each produces {money(unit * 400)} estimated new qualification AI cost. "
         "It excludes import, chat, infrastructure and partial-reuse work; do not divide this by all "
         "leads and call it a measured fresh-lead price.")

    page(doc, 4, "Accounting and cost report interpretation")
    table(doc, ["Work stage", "Typical metered activity", "Attribution"], [
        ["Import / column mapping", "AI mapping when the parser requires it; parsing itself has no model tokens.", "Job level; no arbitrary lead is charged."],
        ["Web research", "Responses model input/output plus reported search requests.", "Lead and job."],
        ["Contact checks", "Syntax, DNS and numbering-plan checks; no separate LLM call in this stage.", "Compute/network overhead outside the AI token bill."],
        ["ICP comparison", "Model interprets supplied evidence and required criteria.", "Lead and job."],
        ["Scoring / outreach", "Numeric scoring is deterministic; permitted-channel outreach can use an LLM.", "Lead and job; no claim of a fixed call count."],
        ["Indexing / query embedding", "Embedding input tokens; Qdrant/storage are separate costs.", "Associated context when available."],
        ["Chat", "Planner, answer generation and retrieval/query embeddings as required.", "Job level when no lead is attributable."],
    ], [1.42, 3.2, 2.25])
    heading(doc, "Completeness rules")
    para(doc, "Every paid attempt is journaled before sending; separate retries have separate IDs. "
         "Interrupted attempts remain started/unreported rather than appearing free. A model's returned "
         "billing identity must match the configured model family to price tokens. A missing search meter "
         "is unknown, not zero. Visible web_search_call items are diagnostic only.")
    para(doc, "The recorder currently prices missing cached-token counts using zero cached tokens, "
         "but marks the attempt incomplete. This provisional component is not a confirmed bill. "
         "Usage is persisted asynchronously, so drilldowns can reflect different instants during processing.")
    para(doc, "In Admin Cost, inspect job → lead → individual attempt, then compare stage totals, "
         "incomplete/unpriced/legacy counts, job-level costs and globally unallocated attempts. "
         "Rate edits affect future snapshots; they do not rewrite historical recorded tariffs.")

    page(doc, 5, "Operating budget and reconciliation")
    table(doc, ["Expense", "In AI estimate?", "How to budget"], [
        ["API and Celery compute", "No", "Instance/container hours × provider rate; include idle and in-flight work."],
        ["Azure SQL", "No", "Selected compute tier, storage, backup retention and provisioned capacity."],
        ["Redis broker / journals", "No", "Memory, persistence, HA tier and replication; retain durable usage data."],
        ["Blob Storage", "No", "Uploaded/archive bytes, operations, retention and data transfer."],
        ["Qdrant", "No", "Cluster tier, vector count/dimension, replicas and storage."],
        ["DNS, networking and logs", "No", "Service charges, egress, monitoring retention and support."],
        ["Tax and FX spread", "No", "Actual invoice jurisdiction, payment terms and effective exchange rate."],
    ], [1.8, 1.1, 3.97])
    heading(doc, "A practical approval checklist")
    para(doc, "1. Confirm Azure tariffs and model deployment names. 2. Run a representative pilot "
         "containing fresh and reused leads. 3. Record per-stage token/search distributions, retries "
         "and incomplete accounting. 4. Multiply expected fresh/reuse/refresh/chat volumes by observed "
         "unit costs. 5. Add fixed infrastructure, tax and an explicit contingency chosen by the budget "
         "owner. 6. Reconcile estimates against Azure billing and investigate unattributed or missing usage.")
    para(doc, "Increasing concurrency (supported limit: 1–30 active leads across workers) primarily "
         "changes spend rate and throughput, not the tariff. Excessive concurrency can increase throttling "
         "and retry overhead. Pausing prevents new qualification starts but does not stop provisioned "
         "infrastructure charges or cancel paid calls already in flight.")
    heading(doc, "Implementation and pricing references")
    for text in [
        "Repository: backend/app/admin/settings.py (reference tariffs); billing.py (metering formulas); "
        "telemetry.py (durable accounting); routes.py (reporting).",
        "Workflow: backend/app/tasks/tasks.py; app/scoring/scoring_engine.py; RELEASE_ADMIN_OPERATIONS.md; RELEASE_JOB_CONTROLS.md.",
        "Azure OpenAI pricing: https://azure.microsoft.com/en-us/pricing/details/azure-openai/",
        "Bing pricing reference configured by the application: https://www.microsoft.com/en-us/bing/apis",
        "These URLs are reference destinations, not evidence of a pricing lookup performed for this document. Confirm current commercial terms separately.",
    ]:
        para(doc, text)
    doc.save(ROOT / "docs/Quality_Lead_Costing_Guide.docx")


def scoring_guide():
    doc = document("Scoring Guide", "Evidence rules, deterministic formulas and worked examples")
    para(doc, f"Policy version: {SCORING_VERSION}. This guide describes the current implementation, "
         "including score caps and contact requirements. It supersedes historical five-factor examples "
         "in older documentation. Authority and data-accuracy columns are inactive legacy storage, not "
         "additional weighted components.")
    heading(doc, "01  Final score and priority")
    table(doc, ["Component", "Weight", "Range", "Purpose"], [
        [name.upper() if name == "icp" else name.title(), f"{weight:.0%}", "0–100",
         {"icp": "Required customer-profile evidence", "contact": "Permitted contact-route availability", "growth": "Sourced, dated demand indicators"}[name]]
        for name, weight in WEIGHTS.items()
    ], [1.3, .75, .85, 3.97])
    para(doc, "Raw score = ICP × 0.60 + Contact × 0.25 + Growth × 0.15. The code then applies "
         "qualification caps, clamps to 0–100 and rounds to one decimal. It never accepts model arithmetic.")
    table(doc, ["Priority", "Final score"], [["High", "80.0–100.0"], ["Medium", "50.0–79.9"], ["Low", "0.0–49.9"]], [1.5, 5.37])
    table(doc, ["Qualification state", "Rule", "Final cap"], [
        ["qualified", "All required ICP checks match and a permitted contact route matches.", "No extra cap"],
        ["needs_review", "No confirmed disqualifier, but required ICP/contact evidence is unresolved.", "79.9"],
        ["disqualified", "At least one required ICP mismatch is established.", "49.0"],
        ["uncontactable", "The contact requirement mismatches (takes precedence over ICP mismatch).", "49.0"],
    ], [1.35, 4.57, .95])
    para(doc, "Qualification state and priority are different: a fully qualified lead without growth "
         "can still be High (85.0); low numerical fit does not automatically prove a mismatch.")

    page(doc, 2, "ICP evidence and points")
    table(doc, ["Field", "What is compared"], [
        [label, {"employee_count": "Researched company size against the requested range.",
                 "location": "Company location against target geography.",
                 "annual_revenue": "Researched revenue against the requested range.",
                 "industry": "Evidence-backed company industry.",
                 "seniority": "Current role/function and required seniority.",
                 "employment": "Current employment at the uploaded company (always required)."}[key]]
        for key, label in FIELDS.items()
    ], [1.65, 5.22])
    table(doc, ["Verdict", "ICP points", "Treatment"], [
        ["match", "1.0", "Required criterion is supported and matched."],
        ["mismatch", "0.0", "Confirmed required contradiction; final score capped at 49."],
        ["unclear / unresolved", "0.5", "Required evidence unresolved; cannot be High until resolved."],
        ["not_specified / optional", "Excluded", "Not counted in the required-criteria denominator."],
    ], [1.8, .9, 4.17])
    para(doc, "ICP score = 100 × total points / required criteria count, rounded to one decimal. "
         "No required criteria yields 0 and a pending assessment. Contact method and growth checks "
         "are excluded from this ICP average; employment is always included.")
    para(doc, "Example: four matches plus one unclear required criterion produce "
         "(4 + 0.5) / 5 × 100 = 90.0. This is not a complete match; even strong contact/growth "
         "cannot override the needs_review cap.")
    heading(doc, "Proof safeguards")
    para(doc, "Confirmed evidence needs a value, a safe HTTP(S) source URL returned by the search "
         "tool, and a supporting excerpt. Unsupported comparisons become unclear. Current employment "
         "matches; former employment mismatches; no reliable current proof is unclear. Additional "
         "mandatory criteria must quote the actual ICP and use allowed supporting sources. "
         "Optional language such as prefer, ideally or nice to have does not create a hard disqualifier.")

    page(doc, 3, "Contact availability")
    para(doc, "The default rule needs at least one email or phone route. Requiring both channels "
         "when the ICP does not ask for both would incorrectly penalize leads; this policy avoids that.")
    table(doc, ["ICP requirement", "Passing channels"], [
        ["Default / email OR phone", "Any available email or phone."],
        ["Email explicitly required", "An available email; a phone alone is insufficient."],
        ["Phone explicitly required", "An available phone; an email alone is insufficient."],
        ["Both explicitly required", "Both email and phone available."],
        ["Verified/deliverable email explicitly required", "Mailbox delivery is not established by these checks; email stays unresolved."],
    ], [2.5, 4.37])
    table(doc, ["Contact outcome", "Score", "Qualification effect"], [
        ["match", "100", "A permitted route is established."],
        ["unclear", "50", "No pass, but a required/default channel has unknown evidence; needs review."],
        ["mismatch", "0", "No permitted route established and no relevant unknown channel; uncontactable cap 49."],
    ], [1.5, .7, 4.67])
    para(doc, "Method states available and unverified count as available routes; unknown does not "
         "establish a pass. Status requires a nonempty channel value. In both-channel mode, the current "
         "implementation treats any needed unknown channel as pending when the overall requirement "
         "does not pass; do not infer that the other channel was proven deliverable.")
    heading(doc, "What the score does and does not establish")
    para(doc, "Email syntax/DNS and phone numbering-plan checks establish limited availability "
         "evidence. They do not establish inbox delivery, phone ownership or a connected call. SMTP "
         "disabled, timeouts and missing delivery evidence do not prove that a mailbox is invalid. "
         "Strict verified-email requirements remain unresolved rather than being falsely certified.")
    para(doc, "Outreach uses a permitted channel only. If no channel matches, the deterministic draft "
         "uses channel none and no opening message. AI may improve wording, but it cannot change "
         "scores, eligibility or the selected contact route.")

    page(doc, 4, "Growth signals and recency")
    table(doc, ["Event type", "Base strength"], [[key.replace("_", " ").title(), value] for key, value in GROWTH_STRENGTH.items()], [4.9, 1.97])
    table(doc, ["Age at research time", "Multiplier", "Example: funding (80)"], [
        ["0–180 days inclusive", "1.00", "80"],
        ["181–365 days inclusive", "0.80", "64"],
        ["366–730 days inclusive", "0.50", "40"],
        ["Future, older than 730 days, or no valid date", "Excluded", "No points"],
    ], [3.25, 1.1, 2.52])
    para(doc, "Each signal must be relevant=true, use a supported event type, include a nonempty "
         "summary and valid date, and cite an allowed tool-returned URL with an excerpt. Duplicate "
         "kind/date/summary entries are removed. The normalizer examines at most 12 input signals.")
    para(doc, "Signal strength = base strength × recency multiplier. Growth score = min(100, "
         "strongest signal + min(10, 5 × (distinct supported event types − 1))). No accepted signals "
         "produces 0. Extra events of the same type do not create a diversity bonus.")
    table(doc, ["Accepted evidence", "Growth score"], [
        ["Recent funding only", "80"],
        ["Recent expansion + recent funding", "min(100, 95 + 5) = 100"],
        ["Recent leadership change only", "45"],
        ["Undated growth summary or no cited event", "0 (unknown demand, not proof of no growth)"],
    ], [4.8, 2.07])

    page(doc, 5, "Worked decisions and audit trail")
    def checks(values):
        return {f"criterion_{index}": {"required": True, "verdict": verdict}
                for index, verdict in enumerate(values)}
    web = {"sources": ["https://example.com/news"], "checked_at": "2026-10-27T00:00:00Z",
           "growth_evidence": {"signals": [{"kind": "hiring", "relevant": True,
               "summary": "Illustrative hiring event", "date": "2026-10-01",
               "sources": [{"url": "https://example.com/news", "excerpt": "Illustrative evidence"}]}]}}
    cases = [
        ("All required criteria match", ["match"] * 5, "available", web),
        ("All match; no sourced growth", ["match"] * 5, "available", {}),
        ("One unclear of five ICP checks", ["match"] * 4 + ["unclear"], "available", web),
        ("One confirmed ICP mismatch", ["match"] * 4 + ["mismatch"], "available", web),
        ("Contact evidence unknown", ["match"] * 5, "unknown", web),
        ("No available contact route", ["match"] * 5, "invalid", web),
    ]
    rows = []
    for name, verdicts, contact_status, growth in cases:
        result = qualification(checks(verdicts), {"methods": [{"kind": "email", "value": "person@example.com", "status": contact_status}]}, growth)
        rows.append([name, f"{result['icp_score']:g}/{result['contact_score']:g}/{result['growth_score']:g}",
                     f"{result['details']['raw_weighted_score']:g}", f"{result['total_score']:g}",
                     result['qualification_status'].replace("_", " ") + " / " + result['priority']])
    table(doc, ["Illustrative case", "ICP/contact/ growth", "Raw", "Final", "State / priority"], rows, [2.05, 1.15, .55, .55, 2.57])
    para(doc, "Examples are synthetic and computed with the actual policy functions by this "
         "document generator. Raw means before qualification caps. Scores are decision support, "
         "not guarantees of buying intent or contact deliverability.")
    heading(doc, "Versioning, reuse and recalculation")
    para(doc, "Results persist the scoring version, component scores, raw weighted score, checks, "
         "reasons and evidence timestamps. Fresh mode skips cross-job reuse but restores this attempt's "
         "successful checkpoints after retries. Resume preserves completed results. Pause does not "
         "change weights or requalify leads. Historical v1 scores are not silently converted by reading "
         "a page; follow the scoring migration/backfill guide for controlled recalculation.")
    heading(doc, "Source of truth")
    table(doc, ["Implementation", "Responsibility"], [
        ["backend/app/scoring/policy.py", "Weights, ICP average, caps, qualification and priority."],
        ["backend/app/validators/contact_policy.py", "Channel requirements and availability points."],
        ["backend/app/validators/evidence.py", "Citation normalization, growth strengths and recency."],
        ["backend/app/scoring/scoring_engine.py", "Deterministic scoring plus optional outreach wording."],
        ["backend/RELEASE_SCORING_V2.md", "Schema/backfill rollout prerequisites."],
    ], [3.5, 3.37])
    doc.save(ROOT / "docs/Quality_Lead_Scoring_Guide.docx")


if __name__ == "__main__":
    cost_guide()
    scoring_guide()
    print("Generated Costing Guide and Scoring Guide in docs/")
